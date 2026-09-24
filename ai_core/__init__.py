from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from PIL import Image
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

try:
    import google.generativeai as genai
except ImportError:  # pragma: no cover
    genai = None

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None


def _coerce_payload(payload: Any) -> dict[str, str]:
    if isinstance(payload, dict):
        return {str(k): str(v) for k, v in payload.items()}
    if hasattr(payload, "model_dump"):
        return _coerce_payload(payload.model_dump())
    if hasattr(payload, "dict"):
        return _coerce_payload(payload.dict())
    return {}


def build_template_document(payload: Any) -> str:
    data = _coerce_payload(payload)
    document_type = data.get("document_type", "Service Agreement")
    party_name = data.get("party_name", "Acme Corp")
    other_party_name = data.get("other_party_name", "Beta LLC")
    role = data.get("role", "Business professional")
    jurisdiction = data.get("jurisdiction", "Delaware, USA")
    effective_date = data.get("effective_date", "2026-10-01")
    deliverables = data.get("deliverables", "services and support")
    purpose = data.get("purpose", "commercial cooperation")
    additional_terms = data.get("additional_terms", "confidentiality, governing law, and good-faith performance")

    return f"""{document_type}

This {document_type} ("Agreement") is entered into as of {effective_date} by and between {party_name} ("Party A") and {other_party_name} ("Party B").

1. Purpose
{purpose}

2. Role and Scope
Party A will engage Party B to provide {role} services and related support. The scope of work includes: {deliverables}.

3. Term and Performance
This Agreement shall commence on {effective_date} and continue until the obligations described herein are fulfilled or terminated in accordance with this Agreement.

4. Confidentiality
Each party shall protect confidential information disclosed by the other party and use such information solely for the purposes contemplated by this Agreement.

5. Governing Law and Dispute Resolution
This Agreement shall be governed by the laws of {jurisdiction}. Any dispute shall first be addressed in good-faith negotiations and, if unresolved, may be subject to mediation or litigation as permitted by law.

6. Additional Terms
{additional_terms}

7. Execution
This Agreement may be executed in counterparts, each of which shall be deemed an original, and all counterparts together shall constitute one and the same instrument.

IN WITNESS WHEREOF, the parties have executed this Agreement as of the Effective Date above.

{party_name}
{other_party_name}
"""


def generate_document(payload: Any) -> str:
    data = _coerce_payload(payload)
    api_key = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
    if genai is not None and api_key:
        try:
            genai.configure(api_key=api_key)
            model = genai.GenerativeModel("gemini-1.5-flash")
            prompt = (
                "Draft a polished professional legal document in clear English. "
                "Keep it practical, structured, and readable. "
                "Return only the final document text without markdown fences. "
                f"Document type: {data.get('document_type', 'Service Agreement')}. "
                f"Party 1: {data.get('party_name', 'Acme Corp')}. "
                f"Party 2: {data.get('other_party_name', 'Beta LLC')}. "
                f"Role: {data.get('role', 'Business professional')}. "
                f"Jurisdiction: {data.get('jurisdiction', 'Delaware, USA')}. "
                f"Effective date: {data.get('effective_date', '2026-10-01')}. "
                f"Deliverables: {data.get('deliverables', 'services and support')}. "
                f"Purpose: {data.get('purpose', 'commercial cooperation')}. "
                f"Additional terms: {data.get('additional_terms', 'confidentiality, governing law, and good-faith performance')}."
            )
            response = model.generate_content(prompt)
            text = getattr(response, "text", None) or str(response)
            if text and text.strip():
                return text.strip()
        except Exception:
            pass
    return build_template_document(data)


def export_document_to_txt(document_text: str, output_path: str) -> str:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(document_text, encoding="utf-8")
    return str(path)


def export_document_to_pdf(document_text: str, output_path: str, company_name: str = "LegalEase", logo_path: str | None = None) -> str:
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    from reportlab.platypus import PageBreak

    doc = SimpleDocTemplate(str(path), pagesize=letter, rightMargin=0.75 * inch, leftMargin=0.75 * inch, topMargin=0.75 * inch, bottomMargin=0.75 * inch)
    styles = getSampleStyleSheet()
    normal = styles["BodyText"]
    normal.fontSize = 10
    normal.leading = 14

    story = []
    story.append(Paragraph(company_name, styles["Title"]))
    story.append(Spacer(1, 0.2 * inch))

    if logo_path and Path(logo_path).exists():
        try:
            img = Image.open(logo_path)
            img_width, img_height = img.size
            max_width = 1.5 * inch
            ratio = max_width / img_width
            max_height = img_height * ratio
            from reportlab.platypus import Image as RLImage
            story.append(RLImage(str(logo_path), width=max_width, height=max_height))
            story.append(Spacer(1, 0.15 * inch))
        except Exception:
            pass

    for paragraph in document_text.split("\n\n"):
        stripped = paragraph.strip()
        if not stripped:
            continue
        story.append(Paragraph(stripped.replace("\n", "<br/>"), normal))
        story.append(Spacer(1, 0.12 * inch))

    doc.build(story)
    return str(path)


def export_document_to_docx(document_text: str, output_path: str, company_name: str = "LegalEase", logo_path: str | None = None) -> str:
    if Document is None:
        raise RuntimeError("python-docx is required to export DOCX files.")

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading(company_name, level=1)

    if logo_path and Path(logo_path).exists():
        try:
            doc.add_picture(str(logo_path), width=200000)
        except Exception:
            pass

    for paragraph in document_text.split("\n\n"):
        cleaned = paragraph.strip()
        if cleaned:
            doc.add_paragraph(cleaned)

    doc.save(str(path))
    return str(path)


__all__ = [
    "build_template_document",
    "generate_document",
    "export_document_to_pdf",
    "export_document_to_docx",
    "export_document_to_txt",
]
